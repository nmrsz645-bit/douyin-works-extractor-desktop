"""抖音博主监控 - Web 前端 (FastAPI)"""
import asyncio
import concurrent.futures
import sys
import logging
from contextlib import asynccontextmanager
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from fastapi import FastAPI, Request, Query, Body, UploadFile, File
from fastapi.responses import HTMLResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from jinja2 import Environment, FileSystemLoader
import uvicorn
import asyncio
import traceback as _tb
import json
import yaml
import threading
import re

from db import (
    init_db, add_creator, remove_creator, rename_creator, list_creators, get_creator,
    upsert_video, add_snapshot, update_last_fetched, update_creator_profile, get_stats, get_db,
    get_all_stats, get_batch_transcripts, get_today_videos, get_today_summary,
    ingest_crawl_results, clear_extraction_results, upsert_comment, list_comments, get_comment_count,
    delete_absent_comments, clear_creator_list, clear_topic_results, upsert_topic_result, list_topic_results,
)
from spider import DouyinSpider, SESSION_DIR
from topic_spider import DouyinTopicSpider, normalise_topic_input
from utils import resolve_secuid, async_resolve_secuid

logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).parent
env = Environment(loader=FileSystemLoader(str(BASE_DIR / "templates")))
env.filters["format_number"] = lambda v: f"{v:,}"
env.filters["format_number_cn"] = lambda v: (
    f"{v/100000000:.1f}亿" if v >= 100000000 else
    f"{v/10000:.1f}万" if v >= 10000 else
    f"{v:,}"
)
env.filters["int"] = int


@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    logger.info("数据库已初始化")
    yield


app = FastAPI(title="抖音博主监控", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
_login_active = False
_extract_state = {"running": False, "current": 0, "total": 0, "found": 0, "creator": "", "message": ""}
_topic_state = {"running": False, "found": 0, "topic": "", "message": ""}
_download_state = {"running": False, "total": 0, "done": 0, "failed": 0, "message": ""}


def render(name: str, **ctx) -> HTMLResponse:
    template = env.get_template(name)
    return HTMLResponse(template.render(**ctx))


def _profile_from_user(user: dict) -> dict:
    """将抖音作者接口中的 user 字段转换为本程序保存的资料格式。"""
    avatar_list = user.get("avatar_medium", {}).get("url_list") or user.get("avatar_thumb", {}).get("url_list") or []
    return {
        "nickname": str(user.get("nickname", "") or ""),
        "avatar_url": avatar_list[0] if avatar_list else "",
        "follower_count": int(user.get("follower_count", 0) or 0),
        "following_count": int(user.get("following_count", 0) or 0),
        "total_likes": int(user.get("total_favorited", 0) or 0),
        "bio": str(user.get("signature", "") or ""),
    }


def _read_creator_profiles_visible(sec_uids: list[str]) -> dict[str, tuple[dict | None, str | None]]:
    """用一个可见浏览器窗口依次读取多位作者资料，不滚动、不固定等待。"""
    unique_sec_uids = list(dict.fromkeys(sec_uid for sec_uid in sec_uids if sec_uid))
    if not unique_sec_uids:
        return {}
    loop = asyncio.ProactorEventLoop() if sys.platform == "win32" else asyncio.new_event_loop()
    asyncio.set_event_loop(loop)

    async def _read_all():
        from config_manager import load_config
        from playwright.async_api import async_playwright

        cfg = load_config()["spider"]
        results: dict[str, tuple[dict | None, str | None]] = {}
        SESSION_DIR.mkdir(parents=True, exist_ok=True)
        async with async_playwright() as p:
            # 资料读取和作品提取使用同一个可见、持久的登录会话。
            context = await p.chromium.launch_persistent_context(
                user_data_dir=str(SESSION_DIR),
                headless=False,
                viewport={"width": cfg["viewport_width"], "height": cfg["viewport_height"]},
                user_agent=cfg["user_agent"],
                locale=cfg["locale"],
            )
            page = context.pages[0] if context.pages else await context.new_page()
            try:
                for sec_uid in unique_sec_uids:
                    profile: dict | None = None
                    profile_ready = asyncio.Event()
                    response_tasks: set[asyncio.Task] = set()

                    async def _consume_profile(response):
                        nonlocal profile
                        if profile is not None or not any(pattern in response.url for pattern in DouyinSpider.PROFILE_PATTERNS):
                            return
                        try:
                            data = await response.json()
                            user = data.get("user", {})
                            if str(user.get("sec_uid", "")) != sec_uid:
                                return
                            profile = _profile_from_user(user)
                            profile_ready.set()
                        except Exception:
                            return

                    def _on_response(response):
                        task = asyncio.create_task(_consume_profile(response))
                        response_tasks.add(task)
                        task.add_done_callback(response_tasks.discard)

                    page.on("response", _on_response)
                    try:
                        await page.goto(f"https://www.douyin.com/user/{sec_uid}", wait_until="domcontentloaded")
                        # 只把 7 秒作为异常网络或验证页的上限；正常接口到达立即进入下一位。
                        await asyncio.wait_for(profile_ready.wait(), timeout=7)
                    except asyncio.TimeoutError:
                        results[sec_uid] = (None, "未收到作者资料(请在弹出浏览器完成登录或验证后重试)")
                    except Exception as exc:
                        results[sec_uid] = (None, f"读取作者主页失败: {exc}")
                    else:
                        results[sec_uid] = (profile, None) if profile else (None, "未收到作者资料")
                    finally:
                        page.remove_listener("response", _on_response)
                        if response_tasks:
                            await asyncio.wait(response_tasks, timeout=1)
            finally:
                await context.close()
        return results

    try:
        return loop.run_until_complete(_read_all())
    finally:
        loop.close()


def _read_creator_profile_once(sec_uid: str):
    """兼容单位作者资料刷新接口。"""
    return _read_creator_profiles_visible([sec_uid]).get(sec_uid, (None, "未收到作者资料"))


def _freshness_level(last_fetched_at):
    if not last_fetched_at:
        return "old"
    try:
        from datetime import datetime, timedelta
        dt = datetime.fromisoformat(last_fetched_at)
        delta = datetime.now() - dt
        if delta < timedelta(hours=6):
            return "fresh"
        elif delta < timedelta(days=1):
            return "stale"
        else:
            return "old"
    except Exception:
        return "old"

env.globals["_freshness_level"] = _freshness_level
env.globals["_freshness_text"] = lambda v: {"fresh": "刚刚", "stale": "今日", "old": "较早"}.get(_freshness_level(v), "较早")


# ─── 页面路由 ──────────────────────────────────────────

@app.get("/", response_class=HTMLResponse)
async def dashboard(request: Request):
    return render("extract.html")
    creators = list_creators()
    all_stats = get_all_stats()
    for c in creators:
        s = all_stats.get(c["id"], {})
        c["total_videos"] = s.get("total_videos", 0)
        c["max_likes"] = s.get("max_likes", 0)
        c["max_comments"] = s.get("max_comments", 0)

    total_likes = 0
    with get_db() as db:
        row = db.execute("""
            SELECT COALESCE(SUM(s.like_count), 0) as total_likes
            FROM snapshots s
            WHERE s.id = (SELECT id FROM snapshots WHERE video_id = s.video_id ORDER BY id DESC LIMIT 1)
        """).fetchone()
        total_likes = row["total_likes"] if row else 0

    today_summary = get_today_summary()
    today_videos = get_today_videos(limit=20)
    for v in today_videos:
        v["create_time_str"] = _fmt_time(v["create_time"])
        v["duration_str"] = _fmt_duration(v["duration_ms"])
    # 批量获取今日视频的转录文本
    if today_videos:
        transcripts = get_batch_transcripts([v["id"] for v in today_videos])
        for v in today_videos:
            v["transcript"] = transcripts.get(v["id"], "")

    return render("dashboard.html",
                  creators=creators,
                  today_summary=today_summary,
                  today_videos=today_videos,
                  summary_stats={
                      "total_creators": len(creators),
                      "total_videos": sum(c.get("total_videos", 0) for c in creators),
                      "total_likes": total_likes,
                  })


@app.get("/creators", response_class=HTMLResponse)
async def creators_page(request: Request):
    return render("creators.html", creators=list_creators())


@app.get("/settings", response_class=HTMLResponse)
def settings_page(request: Request):
    """系统配置管理页面。"""
    from config_manager import load_config
    return render("settings.html", config=load_config())


@app.get("/videos", response_class=HTMLResponse)
async def videos_page(request: Request, creator_id: str = Query(""),
                      search: str | None = Query(None), sort: str = Query("likes"),
                      start_date: str = Query(""), end_date: str = Query(""),
                      page: int = Query(1)):
    cid_int = int(creator_id) if creator_id.strip() else None
    # 校验排序参数（白名单），防止 SQL 注入
    SORT_MAP = {
        "likes": "s.like_count DESC",
        "comments": "s.comment_count DESC",
        "shares": "s.share_count DESC",
        "time": "v.create_time DESC",
    }
    if sort not in SORT_MAP:
        sort = "likes"
    order = SORT_MAP[sort]

    creator = None
    conditions = []
    params = []
    if cid_int:
        conditions.append("v.creator_id = ?")
        params.append(cid_int)
        creator = get_creator(str(cid_int))
    if search:
        conditions.append("v.title LIKE ?")
        params.append(f"%{search}%")
    start_ts, end_ts = _date_range_to_timestamps(start_date, end_date)
    if start_ts is not None:
        conditions.append("v.create_time >= ?")
        params.append(start_ts)
    if end_ts is not None:
        conditions.append("v.create_time < ?")
        params.append(end_ts)
    where_clause = " AND ".join(conditions) if conditions else "1=1"
    offset = (page - 1) * 20

    with get_db() as db:
        rows = db.execute(f"""
            SELECT v.id, v.video_id as vid, v.title, v.cover_url, v.duration_ms,
                   v.create_time, v.hashtags, v.first_seen_at,
                   s.like_count, s.comment_count, s.share_count, s.view_count, s.fetched_at,
                   t.full_text as transcript,
                   c.name as creator_name, c.id as cid
            FROM videos v JOIN snapshots s ON s.video_id = v.id
            JOIN creators c ON c.id = v.creator_id
            LEFT JOIN transcripts t ON t.video_id = v.id
            WHERE {where_clause} AND s.id = (
                SELECT id FROM snapshots WHERE video_id=v.id ORDER BY id DESC LIMIT 1
            )
            ORDER BY {order} LIMIT 20 OFFSET ?
        """, params + [offset]).fetchall()
        total_row = db.execute(
            f"SELECT COUNT(*) as cnt FROM videos v WHERE {where_clause}", params
        ).fetchone()
        total = total_row["cnt"] if total_row else 0

    videos = []
    for r in rows:
        v = dict(r)
        v["create_time_str"] = _fmt_time(v["create_time"])
        v["duration_str"] = _fmt_duration(v["duration_ms"])
        v["hashtags_list"] = json.loads(v["hashtags"]) if v["hashtags"] else []
        videos.append(v)

    return render("videos.html", videos=videos, creator=creator,
                  creators=list_creators(), search=search or "", sort=sort,
                  page=page, total_pages=(total + 19) // 20, total=total,
                  selected_creator=str(cid_int or ""),
                  start_date=start_date, end_date=end_date)


@app.get("/topic", response_class=HTMLResponse)
async def topic_page(request: Request):
    return render("topic.html")


@app.get("/video/{video_db_id}", response_class=HTMLResponse)
async def video_detail(request: Request, video_db_id: int):
    from db import get_db, get_transcript
    with get_db() as db:
        row = db.execute("""
            SELECT v.id, v.video_id as vid, v.title, v.cover_url,
                   v.duration_ms, v.create_time, v.hashtags,
                   s.like_count, s.comment_count, s.share_count,
                   c.name as creator_name, c.id as cid
            FROM videos v
            JOIN snapshots s ON s.video_id = v.id
            JOIN creators c ON c.id = v.creator_id
            WHERE v.id = ?
              AND s.id = (
                  SELECT id FROM snapshots WHERE video_id=v.id ORDER BY id DESC LIMIT 1
              )
        """, (video_db_id,)).fetchone()
        if not row:
            return HTMLResponse("视频不存在", status_code=404)
    v = dict(row)
    v["create_time_str"] = _fmt_time(v["create_time"])
    v["duration_str"] = _fmt_duration(v["duration_ms"])
    v["hashtags_list"] = json.loads(v["hashtags"]) if v["hashtags"] else []
    transcript = get_transcript(video_db_id)
    comment_count = get_comment_count(video_db_id)
    return render("video_detail.html", v=v, transcript=transcript, comment_count=comment_count)
@app.get("/trends/{creator_id}", response_class=HTMLResponse)
async def trends_page(request: Request, creator_id: int):
    creator = get_creator(str(creator_id))
    if not creator:
        return HTMLResponse("博主不存在", status_code=404)
    return render("trends.html", creator=creator)


# ─── API ───────────────────────────────────────────────

@app.post("/api/creators/add")
async def api_add_creator(name: str = Body(""), sec_uid: str = Body("")):
    if not sec_uid:
        return JSONResponse({"error": "sec_uid 不能为空"}, 400)
    sec_uid, err = await async_resolve_secuid(sec_uid)
    if err:
        return JSONResponse({"error": err}, 400)
    if not name:
        name = f"博主_{sec_uid[:8]}"
    cid = add_creator(name, sec_uid)
    return {"id": cid, "name": name}


@app.post("/api/login")
async def api_login():
    """在独立的 Playwright 窗口中扫码；登录态会保存供之后采集使用。"""
    global _login_active
    if _login_active:
        return {"message": "登录窗口已打开，请完成扫码后关闭该窗口"}
    _login_active = True

    def _open_login_window():
        global _login_active
        from spider import SESSION_DIR
        from playwright.async_api import async_playwright

        async def _run():
            async with async_playwright() as p:
                context = await p.chromium.launch_persistent_context(
                    user_data_dir=str(SESSION_DIR), headless=False,
                    viewport={"width": 1440, "height": 900}, locale="zh-CN",
                )
                page = context.pages[0] if context.pages else await context.new_page()
                await page.goto("https://www.douyin.com/", wait_until="domcontentloaded")
                # 登录态保存会即时写入 profile；窗口最多保留十分钟。
                await page.wait_for_timeout(600_000)
                await context.close()
        try:
            asyncio.run(_run())
        except Exception:
            logger.exception("打开登录窗口失败")
        finally:
            _login_active = False

    threading.Thread(target=_open_login_window, daemon=True).start()
    return {"message": "登录窗口已打开，请在抖音页面扫码登录"}


@app.post("/api/creators/batch")
async def api_add_creators_batch(payload: dict = Body(...)):
    """逐行导入主页链接；每行可写成“名称<TAB>链接”。"""
    urls = str(payload.get("urls", ""))
    lines = [line.strip() for line in urls.splitlines() if line.strip()]
    if not lines:
        return JSONResponse({"error": "请至少输入一个主页链接"}, 400)

    added, failed = [], []
    for line in lines:
        parts = line.split("\t", 1)
        name = parts[0].strip() if len(parts) == 2 else ""
        source = parts[-1].strip()
        sec_uid, err = await async_resolve_secuid(source)
        if err:
            failed.append({"input": line, "error": err})
            continue
        try:
            cid = add_creator(name or f"博主_{sec_uid[:8]}", sec_uid)
            added.append({"id": cid, "name": name or f"博主_{sec_uid[:8]}"})
        except Exception as exc:
            failed.append({"input": line, "error": str(exc) or type(exc).__name__})
    return {"added": added, "failed": failed}


@app.get("/api/creators")
async def api_list_creators():
    return {"creators": list_creators()}


@app.delete("/api/creators")
async def api_clear_creator_list():
    """清空作者主页列表和作者作品结果；话题提取结果保留。"""
    return clear_creator_list()


@app.post("/api/creators/add-profile")
async def api_add_creator_profiles(payload: dict = Body(...)):
    """添加主页并读取公开作者资料；不保存任何作品。"""
    raw_urls = str(payload.get("urls", ""))
    lines = [line.strip() for line in raw_urls.splitlines() if line.strip()]
    if not lines:
        return JSONResponse({"error": "请粘贴至少一个作者主页链接"}, 400)
    added, failed, pending = [], [], []
    for line in lines:
        parts = line.split("\t", 1)
        display_name = parts[0].strip() if len(parts) == 2 else ""
        sec_uid, err = await async_resolve_secuid(parts[-1].strip())
        if err:
            failed.append({"input": line, "error": err})
            continue
        cid = add_creator(display_name or f"博主_{sec_uid[:8]}", sec_uid)
        pending.append((cid, sec_uid, display_name, line))

    profile_results = await asyncio.get_running_loop().run_in_executor(
        None, _read_creator_profiles_visible, [item[1] for item in pending]
    )
    for cid, sec_uid, display_name, line in pending:
        profile, read_error = profile_results.get(sec_uid, (None, "未收到作者资料"))
        if profile:
            update_creator_profile(cid, profile)
            if not display_name and profile.get("nickname"):
                rename_creator(cid, profile["nickname"])
        elif read_error:
            failed.append({"input": line, "error": f"主页已添加，但资料读取失败：{read_error}"})
        added.append(get_creator(str(cid)))
    return {"added": added, "failed": failed}


@app.post("/api/creators/{creator_id}/refresh-profile")
async def api_refresh_creator_profile(creator_id: int):
    """只刷新一位作者的公开资料；不会提取或写入作品。"""
    creator = get_creator(str(creator_id))
    if not creator:
        return JSONResponse({"error": "作者不存在"}, 404)

    profile, read_error = await asyncio.get_running_loop().run_in_executor(None, _read_creator_profile_once, creator["sec_uid"])
    if not profile:
        return JSONResponse({"error": read_error or "未收到作者资料，请稍后重试或重新扫码登录"}, 502)
    update_creator_profile(creator_id, profile)
    if profile.get("nickname"):
        rename_creator(creator_id, profile["nickname"])
    return {"creator": get_creator(str(creator_id))}


@app.post("/api/creators/refresh-unread")
async def api_refresh_unread_profiles():
    """依次补读所有尚未获取资料的作者；失败项不会阻塞后续作者。"""
    unread = [creator for creator in list_creators() if not creator.get("nickname")]
    refreshed, failed = 0, []
    profile_results = await asyncio.get_running_loop().run_in_executor(
        None, _read_creator_profiles_visible, [creator["sec_uid"] for creator in unread]
    )
    for creator in unread:
        profile, read_error = profile_results.get(creator["sec_uid"], (None, "未收到作者资料"))
        if not profile:
            failed.append({"id": creator["id"], "name": creator["name"], "error": read_error or "未收到作者资料"})
            continue
        update_creator_profile(creator["id"], profile)
        if profile.get("nickname"):
            rename_creator(creator["id"], profile["nickname"])
        refreshed += 1
    return {"total": len(unread), "refreshed": refreshed, "failed": failed}


@app.delete("/api/creators/{creator_id}")
async def api_remove_creator(creator_id: int):
    remove_creator(str(creator_id))
    return {"ok": True}


@app.put("/api/creators/{creator_id}")
async def api_rename_creator(creator_id: int, name: str = Body(..., embed=True)):
    if not name or not name.strip():
        return JSONResponse({"error": "名称不能为空"}, 400)
    ok = rename_creator(creator_id, name.strip())
    if not ok:
        return JSONResponse({"error": "博主不存在"}, 404)
    return {"ok": True, "name": name.strip()}


@app.post("/api/run")
async def api_run_fetch(creator_id: int = None, creator_ids: list[int] | None = Query(None),
                        start_at: str = Query(""), end_at: str = Query(""), max_per_creator: int | None = Query(None)):
    try:
        start_ts, end_ts = _minute_range_to_timestamps(start_at, end_at)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, 400)
    if creator_ids:
        creators = [get_creator(str(cid)) for cid in creator_ids]
        creators = [c for c in creators if c]
    elif creator_id:
        c = get_creator(str(creator_id))
        if not c:
            return JSONResponse({"error": "博主不存在"}, 404)
        creators = [c]
    else:
        creators = list_creators()
    creators = [c for c in creators if c.get("enabled")]
    if not creators:
        return {"message": "没有启用的博主", "total": 0}

    # 用户选择每次提取重新开始；清除旧作品数据但保留作者资料。
    cleared = clear_extraction_results()
    _extract_state.update({"running": True, "current": 0, "total": len(creators), "found": 0, "creator": "", "message": "准备打开浏览器"})

    def _crawl_one(c: dict) -> tuple:
        loop = asyncio.ProactorEventLoop()
        try:
            # 作品提取必须用可见浏览器进入作者主页并实际滚动，避免无头页面拿不到作品列表。
            batch_count = 0

            def _save_batch(videos: list[dict]):
                """在浏览器每次拿到一页作品时立刻入库，供界面轮询显示。"""
                nonlocal batch_count
                for video in videos:
                    video_db_id = upsert_video(c["id"], video)
                    add_snapshot(video_db_id, video)
                batch_count += len(videos)
                _extract_state["found"] += len(videos)
                _extract_state["message"] = f"已即时显示 {batch_count} 条符合条件的作品"

            spider = DouyinSpider(headless=False, max_scrolls=80, page_load_wait=0, idle_limit=20,
                                  start_ts=start_ts, end_ts=end_ts, on_videos=_save_batch,
                                  max_videos=max_per_creator)
            videos = loop.run_until_complete(spider.fetch(c["sec_uid"], max_retries=1))
            profile = spider.profile.to_dict() if spider.profile else None
            if spider._error:
                return [], None, spider._error
            return batch_count, profile, None
        finally:
            loop.close()

    try:
        total = 0
        failed = []
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
            for index, c in enumerate(creators, start=1):
                _extract_state.update({"current": index, "creator": c["name"], "message": "正在打开主页并向下滑动"})
                batch_count, profile, err = await asyncio.get_running_loop().run_in_executor(
                    pool, _crawl_one, c
                )
                if err:
                    logger.error("抓取 %s 失败: %s", c["name"], err)
                    failed.append({"id": c["id"], "name": c["name"], "error": err})
                    _extract_state["message"] = "未读取到作品，继续下一位作者"
                    continue
                # 作品已在回调中按页写入；这里仅更新作者资料和最后提取时间。
                ingest_crawl_results(c["id"], [], profile)
                total += batch_count
                _extract_state["message"] = f"该作者完成：{batch_count} 条符合条件的作品"

            # 自动转录新视频
            try:
                from transcriber import WHISPER_AVAILABLE
                if WHISPER_AVAILABLE:
                    def _transcribe_all():
                        from transcriber import VideoFetcher, transcribe_pending_videos
                        with VideoFetcher() as fetcher:
                            return transcribe_pending_videos(fetcher)
                    n = await asyncio.get_running_loop().run_in_executor(pool, _transcribe_all)
                    logger.info("自动转录: %d 条", n)
            except ImportError:
                pass

        logger.info("抓取完成: %d 条视频入库，失败 %d 位作者", total, len(failed))
        return {"message": "抓取完成", "total": total, "cleared": cleared, "failed": failed}
    except Exception as e:
        detail = f"{type(e).__name__}: {e}" if str(e) else type(e).__name__
        logger.exception("抓取出错")
        return JSONResponse({"error": f"抓取出错: {detail}"}, 500)
    finally:
        _extract_state["running"] = False


@app.get("/api/extract-status")
async def api_extract_status():
    return _extract_state


@app.post("/api/extract")
async def api_extract(payload: dict = Body(...)):
    """仅对已添加作者，按日期范围或每位作者前 N 条批量提取作品。"""
    start_at = str(payload.get("start_at", ""))
    end_at = str(payload.get("end_at", ""))
    raw_limit = payload.get("max_per_creator")
    max_per_creator = None
    if raw_limit not in (None, ""):
        try:
            max_per_creator = int(raw_limit)
        except (TypeError, ValueError):
            return JSONResponse({"error": "前几条必须是正整数"}, 400)
        if not 1 <= max_per_creator <= 1000:
            return JSONResponse({"error": "每位作者前几条需在 1 到 1000 之间"}, 400)
        # 前几条模式不带日期条件，避免两种筛选交叉造成理解歧义。
        start_at = end_at = ""
    try:
        _minute_range_to_timestamps(start_at, end_at, required=max_per_creator is None)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, 400)
    creator_ids = [c["id"] for c in list_creators() if c.get("enabled")]
    if not creator_ids:
        return JSONResponse({"error": "请先添加作者主页"}, 400)
    result = await api_run_fetch(creator_ids=creator_ids, start_at=start_at, end_at=end_at,
                                 max_per_creator=max_per_creator)
    if isinstance(result, JSONResponse):
        return result
    return result


@app.get("/api/topic-status")
async def api_topic_status():
    return _topic_state


@app.get("/api/topic-results")
async def api_topic_results():
    rows = list_topic_results()
    for row in rows:
        row["publish_time"] = _fmt_time(row["create_time"])
        row["url"] = f"https://www.douyin.com/video/{row['video_id']}"
    return {"rows": rows}


@app.post("/api/topic-extract")
async def api_topic_extract(payload: dict = Body(...)):
    """按日期（可设上限）或最新 N 条提取一个抖音话题的视频。"""
    if _extract_state["running"] or _topic_state["running"]:
        return JSONResponse({"error": "已有提取任务正在运行，请等待完成"}, 409)
    raw_topic = str(payload.get("topic", "")).strip()
    try:
        display_topic, _ = normalise_topic_input(raw_topic)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, 400)
    start_at = str(payload.get("start_at", ""))
    end_at = str(payload.get("end_at", ""))
    raw_top = payload.get("top_count")
    raw_max = payload.get("max_results")
    top_count = max_results = None
    try:
        if raw_top not in (None, ""):
            top_count = int(raw_top)
        if raw_max not in (None, ""):
            max_results = int(raw_max)
    except (TypeError, ValueError):
        return JSONResponse({"error": "提取条数必须是正整数"}, 400)
    if top_count is not None:
        if not 1 <= top_count <= 1000:
            return JSONResponse({"error": "最新前几条需在 1 到 1000 之间"}, 400)
        start_at = end_at = ""
        max_videos = top_count
    else:
        if max_results is None or not 1 <= max_results <= 1000:
            return JSONResponse({"error": "日期模式请填写 1 到 1000 条的提取上限"}, 400)
        max_videos = max_results
    try:
        start_ts, end_ts = _minute_range_to_timestamps(start_at, end_at, required=top_count is None)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, 400)

    cleared = clear_topic_results()
    _topic_state.update({"running": True, "found": 0, "topic": display_topic, "message": "准备打开话题页面"})

    def _crawl_topic() -> tuple[int, str | None]:
        loop = asyncio.ProactorEventLoop() if sys.platform == "win32" else asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            batch_count = 0
            def _save_batch(videos: list[dict]):
                nonlocal batch_count
                for video in videos:
                    upsert_topic_result(display_topic, video)
                batch_count += len(videos)
                _topic_state["found"] += len(videos)
                _topic_state["message"] = f"已即时显示 {batch_count} 条符合条件的话题作品"
            spider = DouyinTopicSpider(raw_topic, start_ts=start_ts, end_ts=end_ts,
                                       max_videos=max_videos, on_videos=_save_batch, headless=False)
            loop.run_until_complete(spider.fetch())
            return batch_count, spider._error
        finally:
            loop.close()

    try:
        _topic_state["message"] = "正在打开话题页并向下滑动"
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
            total, error = await asyncio.get_running_loop().run_in_executor(pool, _crawl_topic)
        if error:
            _topic_state["message"] = error
            return JSONResponse({"error": error, "cleared": cleared}, 502)
        _topic_state["message"] = f"提取完成：{total} 条符合条件的话题作品"
        return {"message": "提取完成", "total": total, "cleared": cleared}
    except Exception as exc:
        logger.exception("话题提取出错")
        return JSONResponse({"error": f"话题提取出错: {type(exc).__name__}: {exc}"}, 500)
    finally:
        _topic_state["running"] = False


def _safe_video_filename(title: str, video_id: str, directory: Path) -> Path:
    """只使用标题做文件名；重复标题自动追加序号，避免覆盖已有文件。"""
    clean = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", (title or "").strip()).rstrip(". ")
    stem = (clean or f"抖音作品_{video_id}")[:120]
    candidate = directory / f"{stem}.mp4"
    index = 2
    while candidate.exists():
        candidate = directory / f"{stem} ({index}).mp4"
        index += 1
    return candidate


def _download_video_file(item: dict, directory: Path) -> Path:
    """下载已由抖音页面返回的可访问播放地址；失败由上层记录并继续下一条。"""
    import requests
    url = str(item.get("video_url") or "")
    if not url:
        raise ValueError("该作品未返回可下载的视频地址")
    output = _safe_video_filename(str(item.get("title") or ""), str(item.get("video_id") or ""), directory)
    temporary = output.with_suffix(".part")
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120 Safari/537.36",
        "Referer": "https://www.douyin.com/",
    }
    with requests.get(url, headers=headers, stream=True, timeout=(15, 90)) as response:
        response.raise_for_status()
        with temporary.open("wb") as handle:
            for chunk in response.iter_content(chunk_size=1024 * 512):
                if chunk:
                    handle.write(chunk)
    temporary.replace(output)
    return output


def _choose_download_directory() -> str:
    import tkinter as tk
    from tkinter import filedialog
    dialog = tk.Tk()
    dialog.withdraw()
    dialog.attributes("-topmost", True)
    path = filedialog.askdirectory(parent=dialog, title="选择视频保存文件夹", mustexist=True)
    dialog.destroy()
    return path


def _run_download_batch(items: list[dict], directory: Path):
    try:
        for item in items:
            try:
                output = _download_video_file(item, directory)
                _download_state["done"] += 1
                _download_state["message"] = f"已保存：{output.name}"
            except Exception as exc:
                _download_state["failed"] += 1
                _download_state["message"] = f"下载失败：{item.get('title') or item.get('video_id')}（{exc}）"
                logger.warning("视频下载失败 %s: %s", item.get("video_id"), exc)
    finally:
        _download_state["running"] = False
        _download_state["message"] = (
            f"下载完成：成功 {_download_state['done']} 条，失败 {_download_state['failed']} 条"
        )


@app.get("/api/download-status")
async def api_download_status():
    return _download_state


@app.post("/api/downloads/start")
async def api_start_downloads(payload: dict = Body(...)):
    """一次选目录后，后台依次下载勾选的视频，支持作者与话题两类结果。"""
    if _download_state["running"]:
        return JSONResponse({"error": "已有视频下载任务正在进行"}, 409)
    source = str(payload.get("source", ""))
    raw_ids = payload.get("ids") or []
    if source not in {"author", "topic"} or not isinstance(raw_ids, list):
        return JSONResponse({"error": "下载参数无效"}, 400)
    identifiers = list(dict.fromkeys(str(value) for value in raw_ids if str(value).strip()))
    if not identifiers:
        return JSONResponse({"error": "请先勾选要下载的作品"}, 400)
    if len(identifiers) > 200:
        return JSONResponse({"error": "一次最多下载 200 条作品"}, 400)
    placeholders = ",".join("?" for _ in identifiers)
    with get_db() as db:
        if source == "author":
            try:
                numeric_ids = [int(value) for value in identifiers]
            except ValueError:
                return JSONResponse({"error": "作者作品标识无效"}, 400)
            placeholders = ",".join("?" for _ in numeric_ids)
            rows = db.execute(
                f"SELECT id, video_id, title, video_url FROM videos WHERE id IN ({placeholders})", numeric_ids
            ).fetchall()
        else:
            rows = db.execute(
                f"SELECT video_id, title, video_url FROM topic_results WHERE video_id IN ({placeholders})", identifiers
            ).fetchall()
    items = [dict(row) for row in rows if row["video_url"]]
    if not items:
        return JSONResponse({"error": "勾选作品暂无可访问的视频地址，请重新提取后再试"}, 400)
    try:
        selected_dir = _choose_download_directory()
    except Exception as exc:
        return JSONResponse({"error": f"选择保存位置失败：{exc}"}, 500)
    if not selected_dir:
        return {"ok": False, "cancelled": True}
    directory = Path(selected_dir)
    _download_state.update({"running": True, "total": len(items), "done": 0, "failed": 0,
                            "message": "准备下载最高可获取画质…"})
    threading.Thread(target=_run_download_batch, args=(items, directory), daemon=True).start()
    return {"ok": True, "total": len(items), "path": str(directory)}


def _topic_xlsx_response():
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    rows = list_topic_results()
    wb = Workbook()
    ws = wb.active
    ws.title = "话题提取结果"
    headers = ["话题名", "作者昵称", "发布时间", "标题", "播放量", "点赞数", "评论总数", "作品链接"]
    ws.append(headers)
    for row in rows:
        ws.append([row["topic"], row["author_name"], _fmt_time(row["create_time"]), row["title"],
                   row["view_count"], row["like_count"], row["comment_count"],
                   f"https://www.douyin.com/video/{row['video_id']}"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, width in {"A": 22, "B": 20, "C": 20, "D": 48, "E": 14, "F": 14, "G": 14, "H": 48}.items():
        ws.column_dimensions[column].width = width
    for row_index in range(2, ws.max_row + 1):
        for column_index in (5, 6, 7):
            ws.cell(row_index, column_index).number_format = "#,##0"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


@app.post("/api/topic-results/export-xlsx/save")
async def api_save_topic_xlsx():
    if not list_topic_results():
        return JSONResponse({"error": "暂无可导出的话题结果"}, 400)
    try:
        import tkinter as tk
        from tkinter import filedialog
        dialog = tk.Tk()
        dialog.withdraw()
        dialog.attributes("-topmost", True)
        output_name = filedialog.asksaveasfilename(
            parent=dialog, title="保存话题提取结果", defaultextension=".xlsx",
            initialfile="抖音话题提取.xlsx", filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        dialog.destroy()
        if not output_name:
            return {"ok": False, "cancelled": True}
        Path(output_name).write_bytes(_topic_xlsx_response())
        return {"ok": True, "path": output_name}
    except Exception as exc:
        logger.exception("话题 XLSX 另存为失败")
        return JSONResponse({"error": f"另存为失败: {exc}"}, 500)


@app.post("/api/import-file")
async def api_import_file(file: UploadFile = File(...)):
    """读取 CSV 或 Excel 的首列，返回可直接粘贴的链接列表。"""
    filename = (file.filename or "").lower()
    content = await file.read()
    try:
        if filename.endswith(".xlsx"):
            from io import BytesIO
            from openpyxl import load_workbook
            ws = load_workbook(BytesIO(content), read_only=True, data_only=True).active
            values = [str(row[0]).strip() for row in ws.iter_rows(values_only=True) if row and row[0]]
        elif filename.endswith(".csv"):
            text = content.decode("utf-8-sig")
            values = [line.split(",")[0].strip().strip('"') for line in text.splitlines() if line.strip()]
        else:
            return JSONResponse({"error": "仅支持 .xlsx 或 .csv 文件"}, 400)
    except Exception as exc:
        return JSONResponse({"error": f"读取文件失败：{exc}"}, 400)
    # 常见表头不作为链接导入。
    values = [v for v in values if "链接" not in v and v.lower() not in {"url", "link"}]
    return {"urls": "\n".join(values), "count": len(values)}


@app.get("/api/results")
async def api_results(start_at: str = Query(""), end_at: str = Query("")):
    try:
        start_ts, end_ts = _minute_range_to_timestamps(start_at, end_at)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, 400)
    conditions, params = [], []
    if start_ts is not None:
        conditions.append("v.create_time >= ?")
        params.append(start_ts)
    if end_ts is not None:
        conditions.append("v.create_time <= ?")
        params.append(end_ts)
    where_clause = " AND ".join(conditions) if conditions else "1=1"
    with get_db() as db:
        rows = db.execute(f"""
            SELECT v.id, c.name AS author, v.title, v.create_time, s.view_count, s.like_count, s.comment_count, v.video_id
            FROM videos v JOIN creators c ON c.id=v.creator_id JOIN snapshots s ON s.video_id=v.id
            WHERE {where_clause} AND s.id=(SELECT id FROM snapshots WHERE video_id=v.id ORDER BY id DESC LIMIT 1)
            ORDER BY c.id ASC, v.create_time DESC
        """, params).fetchall()
    return {"rows": [{"id": r["id"], "author": r["author"], "title": r["title"], "publish_time": _fmt_time(r["create_time"]),
                       "view_count": r["view_count"], "like_count": r["like_count"], "comment_count": r["comment_count"],
                       "url": f"https://www.douyin.com/video/{r['video_id']}"} for r in rows]}


@app.get("/api/trends/{creator_id}")
async def api_trends(creator_id: int):
    with get_db() as db:
        times = db.execute("""
            SELECT DISTINCT fetched_at FROM snapshots s
            JOIN videos v ON v.id=s.video_id WHERE v.creator_id=?
            ORDER BY fetched_at DESC LIMIT 10
        """, (creator_id,)).fetchall()
        times = [t["fetched_at"] for t in reversed(times)]

        top = db.execute("""
            SELECT v.id, v.title, MAX(s.like_count) as m
            FROM videos v JOIN snapshots s ON s.video_id=v.id
            WHERE v.creator_id=? GROUP BY v.id ORDER BY m DESC LIMIT 5
        """, (creator_id,)).fetchall()

        datasets = []
        for tv in top:
            points = []
            for t in times:
                row = db.execute("""
                    SELECT like_count FROM snapshots WHERE video_id=? AND fetched_at<=?
                    ORDER BY fetched_at DESC LIMIT 1
                """, (tv["id"], t)).fetchone()
                points.append(row["like_count"] if row else None)
            datasets.append({"label": (tv["title"] or "")[:20], "data": points})

    return {"labels": times, "datasets": datasets}


@app.get("/api/transcript/{video_db_id}")
async def api_transcript(video_db_id: int):
    """获取视频完整转录文本"""
    from db import get_transcript
    t = get_transcript(video_db_id)
    if not t:
        return JSONResponse({"error": "未转录"}, 404)
    return t


@app.post("/api/transcribe/{video_db_id}")
async def api_transcribe_video(video_db_id: int):
    """触发单个视频语音转录"""
    from db import get_transcript

    # 检查视频是否存在
    with get_db() as db:
        row = db.execute(
            "SELECT v.id, v.video_id, c.name FROM videos v JOIN creators c ON c.id=v.creator_id WHERE v.id=?",
            (video_db_id,)
        ).fetchone()
        if not row:
            return JSONResponse({"error": "视频不存在"}, 404)

    # 已转录则直接返回
    existing = get_transcript(video_db_id)
    if existing:
        return {
            "status": "already_done",
            "full_text": existing.get("full_text", ""),
            "segments": existing.get("segments", []),
        }

    video_id = row["video_id"]

    def _do_transcribe():
        from transcriber import VideoFetcher, process_video
        with VideoFetcher() as fetcher:
            return process_video(fetcher, video_id, video_db_id)

    try:
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
            result = await asyncio.get_running_loop().run_in_executor(pool, _do_transcribe)
        if result:
            return {
                "status": "ok",
                "full_text": result.get("full_text", ""),
                "segments": result.get("segments", []),
            }
        else:
            return {"status": "skipped", "message": "已转录或下载失败"}
    except Exception as e:
        logger.exception("转录失败 video_db_id=%d", video_db_id)
        return JSONResponse({"error": f"转录失败: {e}"}, 500)


@app.post("/api/transcribe-all")
async def api_transcribe_all():
    """批量转录所有未转录视频"""

    with get_db() as db:
        rows = db.execute("""
            SELECT v.id, v.video_id FROM videos v
            LEFT JOIN transcripts t ON t.video_id = v.id
            WHERE t.id IS NULL
            ORDER BY v.id DESC
        """).fetchall()

    if not rows:
        return {"status": "ok", "total": 0, "message": "所有视频已转录"}

    video_list = [(r["id"], r["video_id"]) for r in rows]

    def _do_transcribe_all():
        from transcriber import VideoFetcher, process_video
        results = []
        with VideoFetcher() as fetcher:
            for db_id, vid in video_list:
                try:
                    r = process_video(fetcher, vid, db_id)
                    results.append({"video_db_id": db_id, "status": "ok" if r else "skipped"})
                except Exception as e:
                    results.append({"video_db_id": db_id, "status": "error", "error": str(e)})
        return results

    try:
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
            results = await asyncio.get_running_loop().run_in_executor(pool, _do_transcribe_all)
        ok = sum(1 for r in results if r["status"] == "ok")
        err = sum(1 for r in results if r["status"] == "error")
        skipped = sum(1 for r in results if r["status"] == "skipped")
        return {"status": "ok", "total": len(results), "ok": ok, "error": err, "skipped": skipped}
    except Exception as e:
        logger.exception("批量转录失败")
        return JSONResponse({"error": f"批量转录失败: {e}"}, 500)


@app.get("/api/transcriber-config")
async def api_get_transcriber_config():
    """读取转录配置（模型、设备）"""
    config_path = BASE_DIR.parent / "config.yaml"
    try:
        if config_path.exists():
            with open(config_path, "r", encoding="utf-8") as f:
                cfg = yaml.safe_load(f) or {}
            tc = cfg.get("transcriber", {})
            return {
                "model": tc.get("model", "small"),
                "device": tc.get("device", "cpu"),
            }
    except Exception:
        pass
    return {"model": "small", "device": "cpu"}


@app.put("/api/transcriber-config")
async def api_update_transcriber_config(
    model: str = Body(None),
    device: str = Body(None),
):
    """更新转录配置并写回 config.yaml"""
    config_path = BASE_DIR.parent / "config.yaml"
    if not config_path.exists():
        return JSONResponse({"error": "config.yaml 不存在"}, 500)

    try:
        with open(config_path, "r", encoding="utf-8") as f:
            cfg = yaml.safe_load(f) or {}
    except Exception:
        return JSONResponse({"error": "读取 config.yaml 失败"}, 500)

    tc = cfg.setdefault("transcriber", {})
    updated = {}
    if model is not None:
        tc["model"] = model
        updated["model"] = model
    if device is not None:
        tc["device"] = device
        updated["device"] = device

    try:
        with open(config_path, "w", encoding="utf-8") as f:
            yaml.dump(cfg, f, allow_unicode=True, default_flow_style=False, sort_keys=False)
    except Exception as e:
        return JSONResponse({"error": f"写入 config.yaml 失败: {e}"}, 500)

    logger.info("转录配置已更新: %s", updated)
    return {"status": "ok", "updated": updated}


@app.get("/api/stats")
async def api_stats():
    with get_db() as db:
        tc = db.execute("SELECT COUNT(*) as c FROM creators").fetchone()["c"]
        tv = db.execute("SELECT COUNT(*) as c FROM videos").fetchone()["c"]
        ts = db.execute("SELECT COUNT(*) as c FROM snapshots").fetchone()["c"]
        tl = db.execute("""
            SELECT COALESCE(SUM(s.like_count), 0) as total_likes
            FROM snapshots s
            WHERE s.id = (
                SELECT id FROM snapshots WHERE video_id = s.video_id ORDER BY id DESC LIMIT 1
            )
        """).fetchone()["total_likes"]

    return {"total_creators": tc, "total_videos": tv, "total_snapshots": ts, "total_likes": tl}


@app.get("/api/videos/export")
async def api_export_videos(creator_id: str = Query(""), start_date: str = Query(""),
                            end_date: str = Query(""), start_at: str = Query(""), end_at: str = Query("")):
    """按当前作者/发布日期范围导出每条作品的最新互动数据。"""
    conditions, params = [], []
    if creator_id.strip():
        try:
            conditions.append("v.creator_id = ?")
            params.append(int(creator_id))
        except ValueError:
            return JSONResponse({"error": "作者参数无效"}, 400)
    try:
        start_ts, end_ts = _minute_range_to_timestamps(start_at, end_at)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, 400)
    if start_ts is None and end_ts is None:
        start_ts, end_ts = _date_range_to_timestamps(start_date, end_date)
    if start_ts is not None:
        conditions.append("v.create_time >= ?")
        params.append(start_ts)
    if end_ts is not None:
        conditions.append("v.create_time <= ?" if end_at else "v.create_time < ?")
        params.append(end_ts)
    where_clause = " AND ".join(conditions) if conditions else "1=1"
    with get_db() as db:
        rows = db.execute(f"""
            SELECT c.name AS author, v.title, v.create_time, s.like_count,
                   s.comment_count, s.share_count, s.view_count, v.video_id
            FROM videos v
            JOIN creators c ON c.id = v.creator_id
            JOIN snapshots s ON s.video_id = v.id
            WHERE {where_clause} AND s.id=(
                SELECT id FROM snapshots WHERE video_id=v.id ORDER BY id DESC LIMIT 1
            )
            ORDER BY c.id ASC, v.create_time DESC
        """, params).fetchall()
    import csv
    import io
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["作者", "作品标题", "发布时间", "播放量", "点赞数", "评论总数", "作品链接"])
    for row in rows:
        writer.writerow([
            row["author"], row["title"], _fmt_time(row["create_time"]), row["view_count"], row["like_count"],
            row["comment_count"],
            f"https://www.douyin.com/video/{row['video_id']}",
        ])
    return Response(content="\ufeff" + output.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": "attachment; filename=douyin_videos.csv"})


@app.get("/api/videos/export-xlsx")
async def api_export_videos_xlsx(creator_id: str = Query(""), start_date: str = Query(""),
                                 end_date: str = Query(""), start_at: str = Query(""), end_at: str = Query("")):
    """按当前筛选导出 XLSX，字段与提取结果表完全一致。"""
    conditions, params = [], []
    if creator_id.strip():
        conditions.append("v.creator_id = ?")
        params.append(int(creator_id))
    try:
        start_ts, end_ts = _minute_range_to_timestamps(start_at or start_date, end_at or end_date)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, 400)
    if start_ts is not None:
        conditions.append("v.create_time >= ?")
        params.append(start_ts)
    if end_ts is not None:
        conditions.append("v.create_time <= ?")
        params.append(end_ts)
    where_clause = " AND ".join(conditions) if conditions else "1=1"
    with get_db() as db:
        rows = db.execute(f"""
            SELECT c.name AS author, v.title, v.create_time, s.view_count, s.like_count, s.comment_count, v.video_id
            FROM videos v JOIN creators c ON c.id=v.creator_id JOIN snapshots s ON s.video_id=v.id
            WHERE {where_clause} AND s.id=(SELECT id FROM snapshots WHERE video_id=v.id ORDER BY id DESC LIMIT 1)
            ORDER BY c.id ASC, v.create_time DESC
        """, params).fetchall()

    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "提取结果"
    headers = ["作者", "发布时间", "标题", "播放量", "点赞数", "评论总数", "作品链接"]
    ws.append(headers)
    for row in rows:
        ws.append([
            row["author"], _fmt_time(row["create_time"]), row["title"], row["view_count"],
            row["like_count"], row["comment_count"],
            f"https://www.douyin.com/video/{row['video_id']}",
        ])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, width in {"A": 22, "B": 20, "C": 48, "D": 14, "E": 14, "F": 14, "G": 48}.items():
        ws.column_dimensions[column].width = width
    for row_index in range(2, ws.max_row + 1):
        for column_index in (4, 5, 6):
            ws.cell(row_index, column_index).number_format = "#,##0"

    output = BytesIO()
    wb.save(output)
    filename = "douyin_videos.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/videos/export-xlsx/save")
async def api_save_videos_xlsx(payload: dict = Body(...)):
    """弹出 Windows 另存为窗口，由用户选择路径后写入 XLSX。"""
    start_at = str(payload.get("start_at", ""))
    end_at = str(payload.get("end_at", ""))
    response = await api_export_videos_xlsx(
        creator_id="", start_date="", end_date="", start_at=start_at, end_at=end_at
    )
    if isinstance(response, JSONResponse):
        return response
    try:
        import tkinter as tk
        from tkinter import filedialog
        dialog = tk.Tk()
        dialog.withdraw()
        dialog.attributes("-topmost", True)
        output_name = filedialog.asksaveasfilename(
            parent=dialog,
            title="保存抖音作品提取结果",
            defaultextension=".xlsx",
            initialfile="抖音作品提取.xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        dialog.destroy()
        if not output_name:
            return {"ok": False, "cancelled": True}
        output_path = Path(output_name)
        output_path.write_bytes(response.body)
        return {"ok": True, "path": str(output_path)}
    except Exception as exc:
        logger.exception("XLSX 另存为失败")
        return JSONResponse({"error": f"另存为失败: {exc}"}, 500)


# ─── 配置 API ─────────────────────────────────────────────

@app.get("/api/config")
def api_get_config():
    """返回全部配置（JSON）。"""
    from config_manager import load_config
    return load_config()


@app.post("/api/config")
def api_save_config(data: dict):
    """保存配置到 config.yaml。"""
    from config_manager import save_config, load_config

    def _coerce(v):
        """将字符串值转为正确类型（bool/int/float/str）。"""
        if isinstance(v, str):
            if v.lower() == "true":
                return True
            if v.lower() == "false":
                return False
            try:
                return int(v)
            except (ValueError, TypeError):
                pass
            try:
                return float(v)
            except (ValueError, TypeError):
                pass
        return v

    def _coerce_dict(d: dict) -> dict:
        return {k: _coerce_dict(v) if isinstance(v, dict) else _coerce(v) for k, v in d.items()}

    current = load_config()
    coerced = _coerce_dict(data)
    for section in coerced:
        if section in current and isinstance(current[section], dict):
            current[section].update(coerced[section])

    save_config(current)
    return {"status": "ok"}


# ─── 评论 API ─────────────────────────────────────────────

@app.get("/api/comments/{video_db_id}")
async def api_list_comments(video_db_id: int):
    """返回指定视频的已存评论（JSON）。"""
    rows = list_comments(video_db_id)
    return {"comments": rows, "total": len(rows)}


@app.post("/api/comments/fetch/{video_db_id}")
async def api_fetch_comments(video_db_id: int):
    """触发抓取指定视频的评论。（在独立线程跑 Playwright，避免 Python 3.14 事件循环不兼容）"""
    from comment_spider import CommentSpider

    with get_db() as db:
        video = db.execute(
            "SELECT id, video_id FROM videos WHERE id = ?",
            (video_db_id,),
        ).fetchone()
    if not video:
        return JSONResponse({"error": "视频不存在"}, 404)

    logger.info("开始抓取评论 video_db_id=%s", video_db_id)

    def _sync_fetch(vid: str) -> list[dict]:
        """在独立线程中创建自己的事件循环跑 Playwright。"""
        if sys.platform == "win32":
            asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            spider = CommentSpider(headless=True)
            return loop.run_until_complete(
                spider.fetch_comments(vid, max_pages=3, max_total=30)
            )
        finally:
            loop.close()

    loop = asyncio.get_event_loop()
    try:
        comments = await asyncio.wait_for(
            loop.run_in_executor(None, _sync_fetch, video["video_id"]),
            timeout=60,
        )
    except asyncio.TimeoutError:
        return JSONResponse({"error": "抓取超时"}, 504)
    except Exception as e:
        logger.error("抓取失败: type=%s msg=%r\n%s", type(e).__name__, str(e), _tb.format_exc())
        err_msg = f"{type(e).__name__}: {e}" if str(e) else type(e).__name__
        return JSONResponse({"error": err_msg}, 500)

    saved = 0
    for c in comments:
        upsert_comment(video_db_id, c)
        saved += 1

    active_ids = [c["comment_id"] for c in comments]
    delete_absent_comments(video_db_id, active_ids)

    return {"saved": saved, "total": len(comments)}


def _fmt_time(ts):
    if not ts:
        return "-"
    from datetime import datetime
    return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M")


def _date_range_to_timestamps(start_date: str, end_date: str):
    """返回闭区间起点和开区间终点；空值表示不限制。"""
    from datetime import datetime, timedelta
    try:
        start_ts = int(datetime.strptime(start_date, "%Y-%m-%d").timestamp()) if start_date else None
        end_ts = int((datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)).timestamp()) if end_date else None
    except ValueError:
        return None, None
    return start_ts, end_ts


def _minute_range_to_timestamps(start_at: str, end_at: str, required: bool = False):
    """datetime-local 的分钟范围：起点包含，终点包含该分钟内的全部作品。"""
    from datetime import datetime
    if required and (not start_at or not end_at):
        raise ValueError("请选择开始和结束时间（精确到分钟）")
    try:
        start_ts = int(datetime.strptime(start_at, "%Y-%m-%dT%H:%M").timestamp()) if start_at else None
        end_ts = int(datetime.strptime(end_at, "%Y-%m-%dT%H:%M").timestamp()) + 59 if end_at else None
    except ValueError:
        raise ValueError("时间格式无效，请精确到分钟")
    if start_ts is not None and end_ts is not None and start_ts > end_ts:
        raise ValueError("结束时间不能早于开始时间")
    return start_ts, end_ts


def _fmt_duration(ms):
    if not ms:
        return "-"
    m, s = divmod(ms // 1000, 60)
    return f"{m}:{s:02d}"


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    uvicorn.run("app:app", host="127.0.0.1", port=8080, reload=True)
